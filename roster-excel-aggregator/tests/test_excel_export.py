from io import BytesIO
from openpyxl import load_workbook
from app import excel_export


def test_rows_to_xlsx_roundtrip():
    rows = [
        {"name": "Alice", "date": "1-Aug", "hours": 9.5, "source": "Branch A / August"},
        {"name": "Alice", "date": "3-Aug", "hours": 12.0, "source": "Branch A / August"},
    ]
    data = excel_export.rows_to_xlsx(rows)

    wb = load_workbook(BytesIO(data))
    ws = wb.active
    values = [[cell.value for cell in row] for row in ws.iter_rows()]

    assert values[0] == ["Name", "Date", "Time", "Source"]
    assert values[1] == ["Alice", "1-Aug", 9.5, "Branch A / August"]
    assert values[2] == ["Alice", "3-Aug", 12.0, "Branch A / August"]


def test_rows_to_xlsx_empty_rows_still_has_header():
    data = excel_export.rows_to_xlsx([])
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    values = [[cell.value for cell in row] for row in ws.iter_rows()]
    assert values == [["Name", "Date", "Time", "Source"]]
